"""Executa o root_agent com Runner + sessão in-memory, como no tutorial ADK."""

from __future__ import annotations

import io
import os
import uuid
from dataclasses import dataclass

from google.adk.runners import Runner
from google.adk.sessions import InMemorySessionService
from google.genai import types

from app.agent.agent import CLOUD_FUNCTION_MAP, root_agent
from app.domain.calculators.registry import get_calculator

APP_NAME = "calculadora_gcp"
ALLOWED_TEXT = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
ALLOWED_IMAGE = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}
ALLOWED_FILES = {**ALLOWED_TEXT, **ALLOWED_IMAGE}
MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_FILES = 20


@dataclass
class AgentFillResult:
    filled_services: dict[str, dict[str, float | str]]
    filled_as_is: dict[str, dict[str, float | str]]
    filled_to_be: dict[str, dict[str, float | str]]
    summary: str


@dataclass
class Attachment:
    name: str
    data: bytes


def _extension(filename: str) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot >= 0 else ""


def _extract_docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _extract_csv_text(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _extract_xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    chunks: list[str] = []
    rows_total = 0
    for sheet in workbook.worksheets:
        chunks.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            if any(cell.strip() for cell in cells):
                chunks.append("\t".join(cells))
                rows_total += 1
            if rows_total >= 8000:
                chunks.append("(planilha truncada)")
                return "\n".join(chunks)
    return "\n".join(chunks)


def _extract_xls_text(data: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    chunks: list[str] = []
    rows_total = 0
    for sheet in book.sheets():
        chunks.append(f"# {sheet.name}")
        for row_index in range(sheet.nrows):
            cells = [str(sheet.cell_value(row_index, col)) for col in range(sheet.ncols)]
            if any(cell.strip() for cell in cells):
                chunks.append("\t".join(cells))
                rows_total += 1
            if rows_total >= 8000:
                chunks.append("(planilha truncada)")
                return "\n".join(chunks)
    return "\n".join(chunks)


def _file_parts(filename: str, data: bytes) -> list[types.Part]:
    ext = _extension(filename)
    mime = ALLOWED_FILES.get(ext)
    if mime is None:
        raise ValueError("Use PDF, TXT, CSV, Excel, Word, PNG ou JPG.")
    if ext == ".docx":
        text = _extract_docx_text(data)
        if not text:
            raise ValueError(f"O arquivo Word '{filename}' está vazio.")
        return [types.Part(text=f"Texto extraído de {filename}:\n{text}")]
    if ext == ".csv":
        text = _extract_csv_text(data).strip()
        if not text:
            raise ValueError(f"O CSV '{filename}' está vazio.")
        return [types.Part(text=f"CSV ({filename}):\n{text}")]
    if ext == ".xlsx":
        text = _extract_xlsx_text(data).strip()
        if not text:
            raise ValueError(f"O Excel '{filename}' está vazio.")
        return [types.Part(text=f"Excel ({filename}):\n{text}")]
    if ext == ".xls":
        text = _extract_xls_text(data).strip()
        if not text:
            raise ValueError(f"O Excel '{filename}' está vazio.")
        return [types.Part(text=f"Excel ({filename}):\n{text}")]
    if ext == ".txt":
        return [types.Part(text=f"Arquivo ({filename}):\n{data.decode('utf-8', errors='replace')}")]
    if ext in ALLOWED_IMAGE:
        return [
            types.Part(text=f"Imagem anexada: {filename}"),
            types.Part(inline_data=types.Blob(mime_type=mime, data=data)),
        ]
    return [
        types.Part(text=f"PDF anexado: {filename}"),
        types.Part(inline_data=types.Blob(mime_type=mime, data=data)),
    ]


def _require_api_key() -> None:
    if not os.getenv("GOOGLE_API_KEY") and os.getenv("GOOGLE_GENAI_USE_VERTEXAI", "").upper() not in {
        "TRUE",
        "1",
    }:
        raise RuntimeError(
            "Defina GOOGLE_API_KEY (Google AI Studio) ou autenticação Vertex "
            "(GOOGLE_GENAI_USE_VERTEXAI=TRUE + projeto GCP)."
        )


def _validate_size(name: str, data: bytes) -> None:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"O arquivo '{name}' ultrapassa o limite de 12 MB.")


PROVIDER_LABELS = {
    "gcp": "Google Cloud",
    "azure": "Azure",
    "aws": "AWS",
    "databricks": "Databricks",
    "multicloud": "Multicloud",
}


def _list_call(provider: str) -> str:
    if provider == "multicloud":
        return "list_calculator_services('multicloud')"
    return f"list_calculator_services('{provider}')"


def _task_prompt(
    intent: str,
    source_provider: str,
    target_provider: str,
    notes: str,
    names: list[str],
    want_cloud_compare: bool = False,
) -> str:
    source = PROVIDER_LABELS.get(source_provider, source_provider)
    target = PROVIDER_LABELS.get(target_provider, target_provider)
    extra = f" Notas extras do usuário (respeite preferências de nuvem): {notes}" if notes.strip() else ""
    files = ", ".join(names)
    if intent == "compare":
        return (
            f"Tarefa: COMPARAR AS IS ({source}) com TO-BE ({target}). "
            f"1) {_list_call(source_provider)} e fill_service(..., scenario='as_is') "
            f"reproduzindo a calculadora anexada"
            f"{' em todas as nuvens que aparecerem' if source_provider == 'multicloud' else ''}. "
            f"2) {_list_call(target_provider)} e fill_service(..., scenario='to_be') "
            f"com a proposta"
            f"{' misturando nuvens se as notas pedirem' if target_provider == 'multicloud' else ''}. "
            f"Anexos: {files}.{extra}"
        )
    if intent == "complement":
        return (
            f"Tarefa: COMPLEMENTAR. AS IS = calculadora atual ({source}). "
            f"TO-BE = proposta ({target}). "
            f"1) fill_service scenario='as_is' só com o que o anexo já lista. "
            f"2) fill_service scenario='to_be' com esses itens MAIS gaps, "
            f"{'podendo usar várias nuvens' if target_provider == 'multicloud' else f'só em {target}'}. "
            f"Não lote o catálogo. Anexos: {files}.{extra}"
        )
    if want_cloud_compare:
        return (
            "Tarefa: CONTEXTO com COMPARAÇÃO DE NUVENS. "
            "Monte a proposta por FUNÇÃO e preencha o equivalente em Google, Azure, AWS e Databricks. "
            "Exemplos: Cloud Storage = S3 = ADLS = dbx_storage; BigQuery = Athena/Redshift = Synapse SQL = SQL Warehouse; "
            "Pub/Sub = Kinesis = Event Hubs; Dataflow = Glue = Data Factory = Jobs/DLT; "
            "Dataproc = EMR = Synapse Spark = All-Purpose; Datastream = DMS; Composer = MWAA. "
            f"{CLOUD_FUNCTION_MAP} "
            "Mesmos volumes em cada linha da tabela. Se não houver equivalente no catálogo, omita essa nuvem nessa função. "
            f"Anexos: {files or 'nenhum'}.{extra} "
            "Chame list_calculator_services('multicloud')."
        )
    return (
        f"Tarefa: CONTEXTO. Proposta de solução ({target}). "
        f"Preencha scenario='to_be' com os serviços necessários"
        f"{' na nuvem da proposta (se não citarem nuvem, use Google)' if target_provider == 'multicloud' else f' só em {target}'}. "
        "Não preencha as outras nuvens só para comparar. "
        f"Anexos: {files or 'nenhum'}.{extra} "
        f"Chame {_list_call(target_provider)}."
    )


async def run_fill_agent(
    *,
    files: list[Attachment],
    notes: str = "",
    intent: str = "context",
    source_provider: str = "gcp",
    target_provider: str = "gcp",
    want_cloud_compare: bool = False,
) -> AgentFillResult:
    _require_api_key()
    if not files and not notes.strip():
        raise ValueError("Descreva o projeto ou anexe arquivos.")
    if len(files) > MAX_FILES:
        raise ValueError(f"No máximo {MAX_FILES} anexos por envio.")

    for item in files:
        _validate_size(item.name, item.data)

    names = [item.name for item in files]
    parts: list[types.Part] = [
        types.Part(text=_task_prompt(intent, source_provider, target_provider, notes, names, want_cloud_compare))
    ]
    for item in files:
        parts.extend(_file_parts(item.name, item.data))

    session_service = InMemorySessionService()
    user_id = "calculator-user"
    session_id = str(uuid.uuid4())
    try:
        await session_service.create_session(
            app_name=APP_NAME,
            user_id=user_id,
            session_id=session_id,
            state={
                "source_provider": source_provider,
                "target_provider": target_provider,
                "intent": intent,
            },
        )
    except TypeError:
        await session_service.create_session(
            app_name=APP_NAME, user_id=user_id, session_id=session_id
        )

    runner = Runner(agent=root_agent, app_name=APP_NAME, session_service=session_service)
    content = types.Content(role="user", parts=parts)

    summary = ""
    async for event in runner.run_async(user_id=user_id, session_id=session_id, new_message=content):
        if event.is_final_response() and event.content and event.content.parts:
            texts = [part.text for part in event.content.parts if getattr(part, "text", None)]
            summary = "\n".join(texts).strip()

    session = await session_service.get_session(
        app_name=APP_NAME, user_id=user_id, session_id=session_id
    )
    filled = {}
    filled_as_is_raw = {}
    filled_to_be_raw = {}
    if session is not None:
        filled = session.state.get("filled_services", {}) or {}
        filled_as_is_raw = session.state.get("filled_as_is", {}) or {}
        filled_to_be_raw = session.state.get("filled_to_be", {}) or {}

    def _keep(bucket: dict, provider: str) -> dict[str, dict[str, float | str]]:
        kept: dict[str, dict[str, float | str]] = {}
        for service_id, inputs in bucket.items():
            calculator = get_calculator(service_id)
            if calculator is None:
                continue
            if provider in {"multicloud", "all", "*"} or calculator.definition.provider == provider:
                kept[service_id] = inputs
        return kept

    as_is = _keep(filled_as_is_raw, source_provider)
    to_be = _keep(filled_to_be_raw, target_provider)
    if not to_be:
        to_be = _keep(filled, target_provider)

    return AgentFillResult(
        filled_services=to_be,
        filled_as_is=as_is,
        filled_to_be=to_be,
        summary=summary,
    )
